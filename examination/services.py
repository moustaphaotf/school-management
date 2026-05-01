"""
Pure-Python computations for grading: scale resolution, subject and term
averages, class ranking, mention assignment, and Result persistence.

Also exposes payload-building and PDF rendering helpers used by views and
admin actions to produce single bulletins or a class ZIP.

All averages are returned in the student's native scale (no /4 GPA conversion).
"""
import io
import zipfile
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from typing import Optional

from django.db import transaction
from django.template.loader import render_to_string
from xhtml2pdf import pisa

from academic.models import (
    AllocatedSubject,
    ClassRoom,
    StudentClassEnrollment,
    Subject,
)
from administration.models import Term
from examination.models import (
    ExaminationListHandler,
    GradeScale,
    MarksManagement,
    Result,
)


TWO_PLACES = Decimal("0.01")


def _quantize(value: Decimal) -> Decimal:
    return value.quantize(TWO_PLACES, rounding=ROUND_HALF_UP)


def get_grade_scale(enrollment: StudentClassEnrollment) -> Optional[GradeScale]:
    """Resolve the GradeScale for an enrollment via classroom > class_level >
    grade_level > grade_scale. Returns None if any link is missing."""
    classroom = enrollment.classroom
    class_level = classroom.name  # FK is named 'name' in ClassRoom (legacy)
    if class_level is None:
        return None
    grade_level = class_level.grade_level
    if grade_level is None:
        return None
    return grade_level.grade_scale


def compute_subject_average(
    enrollment: StudentClassEnrollment,
    subject: Subject,
    term: Term,
) -> Optional[Decimal]:
    """Mean of all marks for this enrollment x subject x exams-of-this-term.
    Each mark is rescaled to the student's grade_scale max before averaging.

    Returns None if no marks exist."""
    scale = get_grade_scale(enrollment)
    if scale is None:
        return None
    scale_max = _scale_max(scale)
    if scale_max is None:
        return None

    marks = MarksManagement.objects.filter(
        student=enrollment,
        subject=subject,
        exam_name__term=term,
    ).select_related("exam_name")

    if not marks.exists():
        return None

    total = Decimal("0")
    count = 0
    for m in marks:
        out_of = Decimal(str(m.exam_name.out_of or 0))
        if out_of <= 0:
            continue
        scaled = (Decimal(str(m.points_scored)) / out_of) * scale_max
        total += scaled
        count += 1

    if count == 0:
        return None
    return _quantize(total / Decimal(count))


def compute_term_average(
    enrollment: StudentClassEnrollment,
    term: Term,
) -> Optional[Decimal]:
    """Weighted average of subject averages by their AllocatedSubject coefficient.
    Subjects without any marks are excluded (no impact on the result).

    Returns None if there are no subjects with marks at all."""
    allocations = AllocatedSubject.objects.filter(
        class_room=enrollment.classroom,
        academic_year=enrollment.academic_year,
        term=term,
    ).select_related("subject")

    weighted_sum = Decimal("0")
    coef_sum = Decimal("0")
    for alloc in allocations:
        sub_avg = compute_subject_average(enrollment, alloc.subject, term)
        if sub_avg is None:
            continue
        coef = Decimal(str(alloc.coefficient or 1))
        weighted_sum += sub_avg * coef
        coef_sum += coef

    if coef_sum == 0:
        return None
    return _quantize(weighted_sum / coef_sum)


def assign_mention(average: Decimal, scale: GradeScale) -> str:
    """Look up the GradeScaleRule whose [min, max] range contains the average.
    Returns letter_grade or empty string if no rule matches."""
    if average is None or scale is None:
        return ""
    rule = scale.gradescalerule_set.filter(
        min_grade__lte=average, max_grade__gte=average
    ).first()
    return rule.letter_grade if rule and rule.letter_grade else ""


def compute_class_ranking(classroom: ClassRoom, term: Term) -> list[dict]:
    """Compute the ranking for all enrolled students in a classroom for a term.
    Returns a list of dicts ordered by descending average:
        [{"enrollment": ..., "average": Decimal | None, "rank": int | None,
          "mention": str}]

    Tied averages share a rank; the next rank skips accordingly (1, 1, 3 pattern).
    Students with no average are listed last with rank=None."""
    enrollments = list(
        StudentClassEnrollment.objects.filter(classroom=classroom).select_related(
            "student", "academic_year", "classroom"
        )
    )

    rows = []
    for enr in enrollments:
        avg = compute_term_average(enr, term)
        scale = get_grade_scale(enr)
        rows.append(
            {
                "enrollment": enr,
                "average": avg,
                "rank": None,
                "mention": assign_mention(avg, scale) if avg is not None else "",
            }
        )

    ranked = [r for r in rows if r["average"] is not None]
    unranked = [r for r in rows if r["average"] is None]
    ranked.sort(key=lambda r: r["average"], reverse=True)

    last_avg: Optional[Decimal] = None
    last_rank = 0
    for idx, row in enumerate(ranked, start=1):
        if last_avg is not None and row["average"] == last_avg:
            row["rank"] = last_rank
        else:
            row["rank"] = idx
            last_rank = idx
            last_avg = row["average"]

    return ranked + unranked


@transaction.atomic
def generate_term_results(classroom: ClassRoom, term: Term) -> list[Result]:
    """Compute averages, ranks and mentions for every student in a classroom
    and persist them as Result rows. Updates existing Result rows for the same
    (student, academic_year, term) tuple, otherwise creates new ones.

    Returns the list of saved Result instances."""
    ranking = compute_class_ranking(classroom, term)
    saved: list[Result] = []
    for row in ranking:
        enr: StudentClassEnrollment = row["enrollment"]
        result, _ = Result.objects.update_or_create(
            student=enr.student,
            academic_year=enr.academic_year,
            term=term,
            defaults={
                "average": row["average"],
                "rank": row["rank"],
                "mention": row["mention"],
            },
        )
        saved.append(result)
    return saved


def _scale_max(scale: GradeScale) -> Optional[Decimal]:
    """Return the max_grade across all rules of a scale (e.g. 20 for /20)."""
    rule = scale.gradescalerule_set.order_by("-max_grade").first()
    if rule is None:
        return None
    return Decimal(str(rule.max_grade))


def build_bulletin_payload(
    enrollment: StudentClassEnrollment, term: Term
) -> dict:
    """Compose the bulletin payload (subjects, averages, rank, mention) for
    one (enrollment, term). Used by both JSON and PDF views, and by the ZIP
    bundler. Pure read; does not persist."""
    scale = get_grade_scale(enrollment)
    allocations = AllocatedSubject.objects.filter(
        class_room=enrollment.classroom,
        academic_year=enrollment.academic_year,
        term=term,
    ).select_related("subject")

    subject_lines = []
    for alloc in allocations:
        sub_avg = compute_subject_average(enrollment, alloc.subject, term)
        coef = Decimal(str(alloc.coefficient or 1))
        weighted = (sub_avg * coef) if sub_avg is not None else None
        subject_lines.append(
            {
                "subject": alloc.subject.name,
                "coefficient": coef,
                "average": sub_avg,
                "weighted": weighted,
            }
        )

    average = compute_term_average(enrollment, term)
    mention = assign_mention(average, scale) if average is not None else ""

    ranking = compute_class_ranking(enrollment.classroom, term)
    rank = next(
        (r["rank"] for r in ranking if r["enrollment"].id == enrollment.id), None
    )
    class_size = StudentClassEnrollment.objects.filter(
        classroom=enrollment.classroom
    ).count()

    return {
        "student_id": enrollment.student_id,
        "student_name": (
            f"{enrollment.student.first_name} {enrollment.student.last_name}".strip()
        ),
        "classroom": (
            f"{enrollment.classroom.name.name if enrollment.classroom.name else ''}"
            f" {enrollment.classroom.stream.name if enrollment.classroom.stream else ''}"
        ).strip(),
        "academic_year": enrollment.academic_year.name,
        "term": term.name,
        "scale": scale.name if scale else "",
        "subjects": subject_lines,
        "average": average,
        "rank": rank,
        "class_size": class_size,
        "mention": mention,
    }


def render_bulletin_pdf(
    enrollment: StudentClassEnrollment, term: Term
) -> Optional[bytes]:
    """Render the bulletin HTML template to PDF bytes. Returns None on failure."""
    payload = build_bulletin_payload(enrollment, term)
    html = render_to_string("examination/bulletin.html", {"b": payload})
    buf = io.BytesIO()
    result = pisa.CreatePDF(html, dest=buf)
    if result.err:
        return None
    return buf.getvalue()


def _safe_filename_part(s: str) -> str:
    """Sanitize a string so it's safe to use inside a ZIP entry filename."""
    keep = []
    for ch in s:
        if ch.isalnum() or ch in ("-", "_"):
            keep.append(ch)
        elif ch == " ":
            keep.append("_")
    return "".join(keep) or "x"


def build_marks_template_xlsx(
    exam: ExaminationListHandler, classroom: ClassRoom
) -> bytes:
    """Generate an .xlsx file with one row per enrolled student of the
    classroom and one column per subject allocated to that classroom for
    the exam's term. The first two columns are admission_number and
    student_name (read-only context for the user). Mark columns are empty.

    Used by the directeur/prof to download a pre-filled template,
    fill in the marks offline, then upload via MarksBulkUploadView."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    if exam.term is None:
        raise ValueError("Exam has no term; cannot derive subjects.")

    enrollments = list(
        StudentClassEnrollment.objects.filter(
            classroom=classroom, academic_year=exam.term.academic_year
        )
        .select_related("student")
        .order_by("student__last_name", "student__first_name")
    )
    allocations = list(
        AllocatedSubject.objects.filter(
            class_room=classroom, term=exam.term
        )
        .select_related("subject")
        .order_by("subject__name")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = f"{classroom.id}-{exam.term.name}"[:31]

    headers = ["admission_number", "student_name"] + [
        a.subject.name for a in allocations
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for enr in enrollments:
        row = [
            enr.student.admission_number or "",
            f"{enr.student.last_name or ''} {enr.student.first_name or ''}".strip(),
        ]
        row.extend([None] * len(allocations))
        ws.append(row)

    # Hint row about the max value (informational only, parser ignores it)
    ws.append(
        ["", f"max points: {exam.out_of}"] + [exam.out_of] * len(allocations)
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_marks_xlsx(
    file_obj,
    exam: ExaminationListHandler,
    classroom: ClassRoom,
    teacher,
    apply_changes: bool = True,
) -> dict:
    """Parse an uploaded .xlsx of marks for one (exam, classroom).

    Expected layout: first row = headers ['admission_number',
    'student_name', <subject name>, ...]. Subsequent rows = one student
    each. Empty cells skipped. Final 'hint' row (admission_number empty)
    is silently ignored.

    Validation phase first — collects all errors. If apply_changes is
    True and there are no errors, writes atomically. If errors exist,
    nothing is written.

    Returns: {'created': N, 'updated': N, 'errors': [...], 'total': N}"""
    from openpyxl import load_workbook

    if exam.term is None:
        return {
            "created": 0,
            "updated": 0,
            "errors": [{"row": 0, "msg": "Exam has no term."}],
            "total": 0,
        }

    wb = load_workbook(filename=file_obj, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return {
            "created": 0,
            "updated": 0,
            "errors": [{"row": 0, "msg": "File is empty or has only a header."}],
            "total": 0,
        }

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    if not headers or headers[0].lower().replace(" ", "_") != "admission_number":
        return {
            "created": 0,
            "updated": 0,
            "errors": [
                {"row": 1, "msg": "First column must be 'admission_number'."}
            ],
            "total": 0,
        }

    # Subject columns are everything from index 2 on, with non-empty header
    subject_columns: list[tuple[int, Subject]] = []
    for idx, header in enumerate(headers[2:], start=2):
        if not header:
            continue
        subject = Subject.objects.filter(name__iexact=header).first()
        if subject is None:
            return {
                "created": 0,
                "updated": 0,
                "errors": [
                    {"row": 1, "col": idx + 1, "msg": f"Unknown subject: {header}"}
                ],
                "total": 0,
            }
        subject_columns.append((idx, subject))

    # Each subject must be allocated to this (classroom, term)
    for _, subject in subject_columns:
        if not AllocatedSubject.objects.filter(
            class_room=classroom, term=exam.term, subject=subject
        ).exists():
            return {
                "created": 0,
                "updated": 0,
                "errors": [
                    {
                        "row": 1,
                        "msg": (
                            f"Subject '{subject.name}' is not allocated to "
                            "this classroom for the exam's term."
                        ),
                    }
                ],
                "total": 0,
            }

    errors: list[dict] = []
    pending: list[dict] = []  # accumulated valid (subject, enrollment, points)

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or row[0] in (None, ""):
            continue
        admission_number = str(row[0]).strip()
        # Skip the trailing 'max points' hint row if it lands here somehow
        if admission_number.lower().startswith("max"):
            continue
        try:
            enrollment = StudentClassEnrollment.objects.get(
                classroom=classroom,
                student__admission_number=admission_number,
                academic_year=exam.term.academic_year,
            )
        except StudentClassEnrollment.DoesNotExist:
            errors.append(
                {
                    "row": row_idx,
                    "msg": (
                        f"No enrollment found for admission_number "
                        f"'{admission_number}' in this classroom."
                    ),
                }
            )
            continue

        for col_idx, subject in subject_columns:
            cell = row[col_idx] if col_idx < len(row) else None
            if cell in (None, ""):
                continue
            try:
                points = float(cell)
            except (TypeError, ValueError):
                errors.append(
                    {
                        "row": row_idx,
                        "col": col_idx + 1,
                        "msg": f"Not a number: {cell!r}",
                    }
                )
                continue
            if points < 0 or points > exam.out_of:
                errors.append(
                    {
                        "row": row_idx,
                        "col": col_idx + 1,
                        "msg": (
                            f"{points} out of range [0, {exam.out_of}] "
                            f"for subject {subject.name}"
                        ),
                    }
                )
                continue
            pending.append(
                {
                    "subject": subject,
                    "enrollment": enrollment,
                    "points": points,
                }
            )

    if errors or not apply_changes:
        return {
            "created": 0,
            "updated": 0,
            "errors": errors,
            "total": 0,
        }

    created, updated = 0, 0
    with transaction.atomic():
        for item in pending:
            obj, was_created = MarksManagement.objects.update_or_create(
                exam_name=exam,
                subject=item["subject"],
                student=item["enrollment"],
                defaults={
                    "points_scored": item["points"],
                    "created_by": teacher,
                },
            )
            created += 1 if was_created else 0
            updated += 0 if was_created else 1

    return {
        "created": created,
        "updated": updated,
        "errors": [],
        "total": created + updated,
    }


def render_class_bulletins_zip(
    classroom: ClassRoom, term: Term
) -> tuple[bytes, int, list[str]]:
    """Render a PDF bulletin for each enrolled student in the classroom for
    the term, and bundle them into an in-memory ZIP. Returns:
        (zip_bytes, success_count, failed_student_names)"""
    enrollments = (
        StudentClassEnrollment.objects.filter(classroom=classroom)
        .select_related("student", "academic_year", "classroom")
        .order_by("student__last_name", "student__first_name")
    )

    buf = io.BytesIO()
    failed: list[str] = []
    success = 0
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for enr in enrollments:
            pdf = render_bulletin_pdf(enr, term)
            student_name = (
                f"{enr.student.last_name} {enr.student.first_name}".strip()
            )
            if pdf is None:
                failed.append(student_name)
                continue
            filename = (
                f"bulletin_{_safe_filename_part(enr.student.last_name or '')}_"
                f"{_safe_filename_part(enr.student.first_name or '')}_"
                f"{_safe_filename_part(term.name)}.pdf"
            )
            zf.writestr(filename, pdf)
            success += 1

    return buf.getvalue(), success, failed
